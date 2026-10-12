import re
import tempfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import fitz
from docx import Document
from fastapi import HTTPException, status

from resume_job_matcher.core.config import Settings
from resume_job_matcher.domain.resume_facts import FactValue, SectionFactValue, SkillFactValue
from resume_job_matcher.domain.skill_dictionary import SkillDictionary
from resume_job_matcher.domain.technologies import find_technology_mentions

PARSER_VERSION = "deterministic-facts-v3"
SUPPORTED_FILE_SUFFIXES = {".pdf": "pdf", ".docx": "docx", ".txt": "text"}
SECTION_TYPES = {
    "个人信息": "contact",
    "个人简介": "summary",
    "自我介绍": "summary",
    "专业技能": "skills",
    "技术栈": "skills",
    "技能": "skills",
    "工作经历": "work_experience",
    "工作经验": "work_experience",
    "项目经历": "project_experience",
    "项目经验": "project_experience",
    "教育经历": "education",
    "教育背景": "education",
}


@dataclass(frozen=True)
class SourceLine:
    text: str
    location: str


@dataclass(frozen=True)
class ParsedFact:
    fact_type: str
    value: FactValue
    source_text: str
    source_location: str


class ResumeImportError(HTTPException):
    def __init__(self, detail: str) -> None:
        super().__init__(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail=detail)


def determine_file_format(filename: str | None) -> str:
    if filename is None:
        raise ResumeImportError("A filename is required for uploaded resumes")
    file_format = SUPPORTED_FILE_SUFFIXES.get(Path(filename).suffix.lower())
    if file_format is None:
        raise ResumeImportError("Only PDF, DOCX, and UTF-8 text files are supported")
    return file_format


@contextmanager
def temporary_raw_file(settings: Settings, content: bytes, suffix: str) -> Iterator[Path]:
    """Contain a raw upload to one directory that is removed on every path."""

    settings.upload_temp_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="upload-", dir=settings.upload_temp_dir) as directory:
        file_path = Path(directory) / f"source{suffix}"
        file_path.write_bytes(content)
        yield file_path


def extract_source_lines(file_path: Path, file_format: str) -> list[SourceLine]:
    if file_format == "pdf":
        return extract_pdf_lines(file_path)
    if file_format == "docx":
        return extract_docx_lines(file_path)
    return extract_text_lines(file_path)


def extract_pdf_lines(file_path: Path) -> list[SourceLine]:
    if not file_path.read_bytes().startswith(b"%PDF-"):
        raise ResumeImportError("The uploaded PDF is invalid")
    try:
        document = fitz.open(file_path)
    except fitz.FileDataError as exc:
        raise ResumeImportError("The uploaded PDF cannot be read") from exc
    try:
        lines: list[SourceLine] = []
        for page_index, page in enumerate(document, start=1):
            for line_index, line in enumerate(page.get_text("text").splitlines(), start=1):
                normalized = line.strip()
                if normalized:
                    lines.append(SourceLine(normalized, f"page:{page_index},line:{line_index}"))
    finally:
        document.close()
    if len("".join(line.text for line in lines)) < 80:
        raise ResumeImportError(
            "Scanned or textless PDFs are not supported; use DOCX or paste text"
        )
    return lines


def extract_docx_lines(file_path: Path) -> list[SourceLine]:
    try:
        with ZipFile(file_path) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ResumeImportError("The uploaded DOCX is invalid")
        document = Document(str(file_path))
    except (BadZipFile, ValueError) as exc:
        raise ResumeImportError("The uploaded DOCX cannot be read") from exc

    lines = [
        SourceLine(paragraph.text.strip(), f"paragraph:{index}")
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                cell_text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )
                if cell_text:
                    lines.append(
                        SourceLine(
                            cell_text, f"table:{table_index},row:{row_index},cell:{cell_index}"
                        )
                    )
    return lines


def extract_text_lines(file_path: Path) -> list[SourceLine]:
    try:
        text = file_path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ResumeImportError("Text files must be UTF-8 encoded") from exc
    return [
        SourceLine(line.strip(), f"line:{index}")
        for index, line in enumerate(text.splitlines(), start=1)
        if line.strip()
    ]


def parse_pasted_text(content: str) -> list[SourceLine]:
    return [
        SourceLine(line.strip(), f"line:{index}")
        for index, line in enumerate(content.splitlines(), start=1)
        if line.strip()
    ]


def build_facts(
    lines: list[SourceLine], dictionary: SkillDictionary | None = None
) -> list[ParsedFact]:
    if not lines:
        raise ResumeImportError("The resume does not contain extractable text")

    sections: list[tuple[str, str, list[SourceLine]]] = []
    current_title = "未分段内容"
    current_type = "other"
    current_lines: list[SourceLine] = []
    for line in lines:
        heading = classify_heading(line.text)
        if heading is not None:
            if current_lines:
                sections.append((current_title, current_type, current_lines))
            current_title, current_type = heading
            current_lines = [line]
        else:
            current_lines.append(line)
    if current_lines:
        sections.append((current_title, current_type, current_lines))

    facts: list[ParsedFact] = []
    for title, fact_type, section_lines in sections:
        source_text = "\n".join(line.text for line in section_lines)
        source_location = f"{section_lines[0].location}..{section_lines[-1].location}"
        facts.append(
            ParsedFact(
                fact_type=fact_type,
                value=SectionFactValue(title=title, content=source_text),
                source_text=source_text,
                source_location=source_location,
            )
        )
        if fact_type == "skills":
            facts.extend(extract_skill_facts(source_text, source_location, dictionary))
    return facts


def extract_skill_facts(
    source_text: str, source_location: str, dictionary: SkillDictionary | None = None
) -> list[ParsedFact]:
    facts: list[ParsedFact] = []
    for mention in find_technology_mentions(source_text, dictionary):
        facts.append(
            ParsedFact(
                fact_type="skill",
                value=SkillFactValue(name=mention.name),
                source_text=mention.source_text,
                source_location=source_location,
            )
        )
    return facts


def classify_heading(text: str) -> tuple[str, str] | None:
    normalized = re.sub(r"[\s【】\[\]#：:·.、0-9一二三四五六七八九十]+", "", text)
    if len(normalized) > 12:
        return None
    fact_type = SECTION_TYPES.get(normalized)
    if fact_type is None:
        return None
    return normalized, fact_type
