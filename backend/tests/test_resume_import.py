from io import BytesIO
from pathlib import Path
from typing import Any, cast

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from resume_job_matcher.api.app import app
from resume_job_matcher.application.auth_service import create_invitation
from resume_job_matcher.application.resume_import_service import temporary_raw_file
from resume_job_matcher.core.config import get_settings
from resume_job_matcher.infrastructure.db import SessionLocal


def create_invite() -> str:
    with SessionLocal.begin() as database_session:
        _, code = create_invitation(
            database_session,
            get_settings(),
            created_by=None,
            max_uses=1,
            expires_in_days=1,
        )
        return code


def authenticated_client(email: str) -> TestClient:
    test_client = TestClient(app)
    response = test_client.post(
        "/api/v1/auth/register",
        json={
            "email": email,
            "password": "correct-password",
            "invitation_code": create_invite(),
            "is_adult": True,
        },
    )
    assert response.status_code == 201
    return test_client


def csrf_headers(client: TestClient) -> dict[str, str]:
    csrf_token = client.cookies.get("rjm_csrf")
    assert csrf_token is not None
    return {"X-CSRF-Token": csrf_token}


def import_text(client: TestClient) -> dict[str, Any]:
    resume_text = "\n".join(
        [
            "张三",
            "个人简介",
            "三年 Python 后端开发经验",
            "专业技能",
            "Python、FastAPI、PostgreSQL",
            "项目经历",
            "负责内部数据平台开发",
        ]
    )
    response = client.post(
        "/api/v1/resumes/import",
        data={"text_content": resume_text},
        headers=csrf_headers(client),
    )
    assert response.status_code == 201
    return cast(dict[str, Any], response.json())


def test_text_import_requires_fact_review_and_enforces_ownership() -> None:
    owner = authenticated_client("owner@example.test")
    imported = import_text(owner)
    profile_id = str(imported["id"])
    facts = cast(list[dict[str, Any]], imported["facts"])
    assert imported["status"] == "pending_confirmation"
    assert {fact["fact_type"] for fact in facts} >= {"summary", "skills", "project_experience"}
    skill_facts = [fact for fact in facts if fact["fact_type"] == "skill"]
    assert {fact["value"]["name"] for fact in skill_facts} == {"Python", "FastAPI", "PostgreSQL"}
    assert {fact["source_text"] for fact in skill_facts} == {"Python", "FastAPI", "PostgreSQL"}
    assert owner.get("/api/v1/resumes").json()[0]["id"] == profile_id

    assert (
        owner.post(f"/api/v1/resumes/{profile_id}/confirm", headers=csrf_headers(owner)).status_code
        == 409
    )

    other_user = authenticated_client("other@example.test")
    assert other_user.get(f"/api/v1/resumes/{profile_id}").status_code == 404
    assert other_user.get("/api/v1/resumes").json() == []

    for fact in facts:
        response = owner.patch(
            f"/api/v1/resumes/{profile_id}/facts/{fact['id']}",
            json={"user_status": "confirmed"},
            headers=csrf_headers(owner),
        )
        assert response.status_code == 200
    confirmed = owner.post(f"/api/v1/resumes/{profile_id}/confirm", headers=csrf_headers(owner))
    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "confirmed"


def test_docx_and_text_pdf_imports_are_supported() -> None:
    client = authenticated_client("formats@example.test")
    docx_buffer = BytesIO()
    document = Document()
    document.add_paragraph("专业技能")
    document.add_paragraph("Python、FastAPI")
    document.save(docx_buffer)
    docx_response = client.post(
        "/api/v1/resumes/import",
        files={
            "file": (
                "resume.docx",
                docx_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=csrf_headers(client),
    )
    assert docx_response.status_code == 201
    assert docx_response.json()["source_format"] == "docx"

    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "Skills: Python FastAPI PostgreSQL " * 8)
    pdf_response = client.post(
        "/api/v1/resumes/import",
        files={"file": ("resume.pdf", pdf_document.tobytes(), "application/pdf")},
        headers=csrf_headers(client),
    )
    pdf_document.close()
    assert pdf_response.status_code == 201
    assert pdf_response.json()["source_format"] == "pdf"


def test_scanned_pdf_is_rejected() -> None:
    client = authenticated_client("scan@example.test")
    pdf_document = fitz.open()
    pdf_document.new_page()
    response = client.post(
        "/api/v1/resumes/import",
        files={"file": ("scan.pdf", pdf_document.tobytes(), "application/pdf")},
        headers=csrf_headers(client),
    )
    pdf_document.close()
    assert response.status_code == 422
    assert "Scanned" in response.json()["detail"]


def test_temporary_raw_file_is_removed_after_parser_failure(tmp_path: Path) -> None:
    settings = get_settings().model_copy(update={"upload_temp_dir": tmp_path})
    with pytest.raises(RuntimeError):
        with temporary_raw_file(settings, b"synthetic content", ".txt") as raw_file:
            assert raw_file.exists()
            raise RuntimeError("synthetic parser failure")
    assert list(tmp_path.iterdir()) == []
